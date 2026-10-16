"""Excel de trazabilidad desde MySQL: referencias de fotos, nunca sus binarios."""
from datetime import timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from sqlalchemy import select

from app.config import BASE_DIR
from app.models import AppArchivo, AppEvidencia, AppIngreso, AppUnidadLogistica
from app.services import app_modelo as m


def generar(db, codigo):
    r=m.obtener(db,codigo)
    ingreso=db.get(AppIngreso,codigo)
    unidades={u.id:u for u in db.scalars(select(AppUnidadLogistica).where(AppUnidadLogistica.recepcion==codigo))}
    libro=Workbook();hoja=libro.active;hoja.title='EVIDENCIAS'
    logo=BASE_DIR/'firebase/public/logo-itsa.png'
    if logo.is_file():
        img=Image(str(logo));img.width=132;img.height=44;hoja.add_image(img,'A1')
    hoja['D1']='TRAZABILIDAD DE EVIDENCIAS · ITSANET';hoja['D1'].font=Font(bold=True,size=14,color='1D4ED8')
    hoja['D2']=f'{r.codigo} · {r.cliente}';hoja['D2'].data_type='s';hoja['D3']='Origen: base de datos del servidor'
    columnas=['CLIENTE','MARCA','RECEPCION','SKU','CLASIFICACION','TIPO','PROTOCOLO','BULTO',
        'CODIGO BULTO','CAJA MADRE','CODIGO CAJA','PACKING LIST','FECHA ECUADOR','USUARIO','ARCHIVO ID','ENLACE DRIVE']
    for i,nombre in enumerate(columnas,1):
        c=hoja.cell(5,i,nombre);c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='1D4ED8')
    filas=db.execute(select(AppEvidencia,AppArchivo).join(AppArchivo,AppArchivo.id==AppEvidencia.archivo_id)
        .where(AppEvidencia.recepcion==codigo,AppEvidencia.eliminado_en.is_(None))
        .order_by(AppEvidencia.sku,AppEvidencia.creado_en,AppEvidencia.id))
    for e,a in filas:
        u=unidades.get(e.unidad_id);b=unidades.get(u.padre_id) if u and u.padre_id else u
        fecha=e.creado_en if e.creado_en.tzinfo else e.creado_en.replace(tzinfo=timezone.utc)
        valores=[r.cliente,(ingreso.datos if ingreso else {}).get('marca',''),codigo,e.sku,e.categoria,e.tipo,e.protocolo,
            b.numero if b else '',b.codigo if b else '',u.numero if u and u.tipo=='CAJA_MADRE' else '',
            u.codigo if u and u.tipo=='CAJA_MADRE' else '',e.origen,fecha.astimezone(ZoneInfo('America/Guayaquil')).strftime('%d/%m/%Y %H:%M:%S'),
            e.usuario,a.id,a.drive_url or 'Pendiente de copia a Drive']
        hoja.append(valores)
        for c in hoja[hoja.max_row]:
            if isinstance(c.value,str):c.data_type='s'
    hoja.freeze_panes='A6';hoja.auto_filter.ref=f'A5:P{max(5,hoja.max_row)}'
    for columna in hoja.columns:hoja.column_dimensions[columna[0].column_letter].width=24
    hoja.column_dimensions['P'].width=60
    salida=BytesIO();libro.save(salida);return salida.getvalue()
