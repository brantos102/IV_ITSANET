"""Importacion del packing list (.xlsx) y del maestro de productos Depot.

El importador detecta la fila de encabezado y mapea las columnas por nombre,
de modo que tolera packing lists con columnas en distinto orden o con
columnas de validacion ya presentes (que se ignoran: la validacion la hace
la app, no el archivo).
"""
from __future__ import annotations

import hashlib
import json
from dataclasses import dataclass, field

from openpyxl import load_workbook

from app.services.normalize import a_numero, limpiar, normalizar_sku

# Alias de encabezado -> campo canonico. Se compara en mayusculas sin acentos.
ALIAS = {
    "ITEM": "sku", "SKU": "sku", "CODIGO": "sku", "PARTE": "sku",
    "NUMERO DE PARTE": "sku", "PART NUMBER": "sku", "PRODUCTO": "sku",
    "DESCRIPCION": "descripcion", "DESCRIPTION": "descripcion", "DETALLE": "descripcion",
    "REFERENCIAPEDIDO": "referencia_pedido", "REFERENCIA PEDIDO": "referencia_pedido",
    "REFERENCIA": "referencia_pedido", "PEDIDO": "referencia_pedido",
    "FACTURA": "factura", "INVOICE": "factura",
    "CANTIDAD": "cantidad", "QTY": "cantidad", "QUANTITY": "cantidad", "CANT": "cantidad",
    "PRECIOUNIT": "precio_unit", "PRECIO UNIT": "precio_unit", "PRECIO UNITARIO": "precio_unit",
    "PRECIOTOTAL": "precio_total", "PRECIO TOTAL": "precio_total",
    "NUMERODEPIEZAREEMPLAZO": "num_pieza_reemplazo",
    "NUMERO DE PIEZA REEMPLAZO": "num_pieza_reemplazo",
    "EMBARQUE": "embarque", "SHIPMENT": "embarque", "BL": "embarque",
}
# Columnas de resultado que la app genera: si vienen en el archivo se descartan.
COLUMNAS_RESULTADO = {
    "CANTIDAD FISICA", "CANTIDADFISICA", "VALIDACION", "PALLET",
    "CODIGO BARRA", "CODIGOBARRA", "COD/PROD", "CODPROD", "OBSERVACION", "OK",
}

OBLIGATORIAS = {"sku", "cantidad"}


@dataclass
class FilaPacking:
    linea: int
    sku: str
    descripcion: str = ""
    referencia_pedido: str = ""
    factura: str = ""
    cantidad: float = 0.0
    precio_unit: float | None = None
    precio_total: float | None = None
    num_pieza_reemplazo: str = ""
    embarque: str = ""
    extra: dict = field(default_factory=dict)


@dataclass
class ResultadoImport:
    filas: list[FilaPacking]
    hoja: str
    encabezados: list[str]
    descartadas: int
    hash_archivo: str
    advertencias: list[str] = field(default_factory=list)

    @property
    def total_lineas(self) -> int:
        return len(self.filas)

    @property
    def total_unidades(self) -> float:
        return sum(f.cantidad for f in self.filas)

    @property
    def skus_unicos(self) -> int:
        return len({f.sku for f in self.filas})

    @property
    def embarques(self) -> list[str]:
        return sorted({f.embarque for f in self.filas if f.embarque})


def _clave(texto) -> str:
    from app.services.normalize import sin_acentos

    return sin_acentos(limpiar(texto)).upper()


def _mapear_encabezado(fila) -> dict[int, str]:
    """Devuelve {indice_columna: campo}. Columnas no reconocidas van a `extra`."""
    mapa: dict[int, str] = {}
    for i, celda in enumerate(fila):
        k = _clave(celda)
        if not k or k in COLUMNAS_RESULTADO:
            continue
        campo = ALIAS.get(k) or ALIAS.get(k.replace(" ", ""))
        mapa[i] = campo or f"extra:{k[:40]}"
    return mapa


def _buscar_encabezado(ws, max_scan: int = 25) -> tuple[int, dict[int, str]]:
    """Localiza la fila de encabezados: la que cubra sku + cantidad."""
    mejor: tuple[int, dict[int, str]] | None = None
    for idx, fila in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan, values_only=True), start=1
    ):
        mapa = _mapear_encabezado(fila)
        campos = {v for v in mapa.values() if not v.startswith("extra:")}
        if OBLIGATORIAS <= campos:
            return idx, mapa
        if mejor is None and "sku" in campos:
            mejor = (idx, mapa)
    if mejor:
        return mejor
    raise ValueError(
        "No se encontro la fila de encabezados. Se requieren al menos las "
        "columnas ITEM (o SKU) y CANTIDAD."
    )


def hash_archivo(ruta) -> str:
    h = hashlib.sha256()
    with open(ruta, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def leer_packing_list(ruta, hoja: str | None = None) -> ResultadoImport:
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else _elegir_hoja(wb)
    fila_hdr, mapa = _buscar_encabezado(ws)

    filas: list[FilaPacking] = []
    descartadas = 0
    advertencias: list[str] = []
    linea = 0

    for raw in ws.iter_rows(min_row=fila_hdr + 1, values_only=True):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        datos: dict = {}
        extra: dict = {}
        for i, campo in mapa.items():
            if i >= len(raw):
                continue
            valor = raw[i]
            if campo.startswith("extra:"):
                v = limpiar(valor)
                if v:
                    extra[campo[6:]] = v
            else:
                datos[campo] = valor

        sku = normalizar_sku(datos.get("sku"))
        if not sku:
            descartadas += 1
            continue

        linea += 1
        filas.append(
            FilaPacking(
                linea=linea,
                sku=sku,
                descripcion=limpiar(datos.get("descripcion")),
                referencia_pedido=limpiar(datos.get("referencia_pedido")),
                factura=limpiar(datos.get("factura")),
                cantidad=a_numero(datos.get("cantidad")),
                precio_unit=a_numero(datos.get("precio_unit"), 0.0) or None,
                precio_total=a_numero(datos.get("precio_total"), 0.0) or None,
                num_pieza_reemplazo=limpiar(datos.get("num_pieza_reemplazo")),
                embarque=limpiar(datos.get("embarque")),
                extra=extra,
            )
        )

    wb.close()
    if not filas:
        raise ValueError("El archivo no contiene lineas con SKU validas.")
    sin_cantidad = sum(1 for f in filas if f.cantidad <= 0)
    if sin_cantidad:
        advertencias.append(f"{sin_cantidad} linea(s) con cantidad 0 o vacia.")

    return ResultadoImport(
        filas=filas,
        hoja=ws.title,
        encabezados=[v for v in mapa.values()],
        descartadas=descartadas,
        hash_archivo=hash_archivo(ruta),
        advertencias=advertencias,
    )


def _elegir_hoja(wb):
    """Primera hoja visible cuyo encabezado sea reconocible."""
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        try:
            _buscar_encabezado(ws)
            return ws
        except ValueError:
            continue
    return wb.worksheets[0]


# --------------------------------------------------------------------------- #
# Maestro de productos (export del sistema Depot)
# --------------------------------------------------------------------------- #
MAESTRO_ALIAS = {
    "producto_id": "sku", "codigo_producto": "codigo_alternativo",
    "descripcion": "descripcion", "codigo_barra": "ean", "unidad_id": "unidad",
    "familia_id": "familia", "cliente_id": "cliente", "inactivo": "inactivo",
    "CODIGO_ALTERNATIVO": "codigo_alternativo",
}


def leer_maestro(ruta, hoja: str | None = None):
    """Itera el maestro de productos y rinde dicts listos para `Producto`."""
    wb = load_workbook(ruta, data_only=True, read_only=True)
    ws = wb[hoja] if hoja else (wb["MAESTRO"] if "MAESTRO" in wb.sheetnames else wb.worksheets[0])
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    mapa = {}
    for i, h in enumerate(hdr):
        k = limpiar(h)
        if k in MAESTRO_ALIAS:
            mapa[i] = MAESTRO_ALIAS[k]
        elif k.lower() in MAESTRO_ALIAS:
            mapa[i] = MAESTRO_ALIAS[k.lower()]

    if "sku" not in mapa.values():
        wb.close()
        raise ValueError("El maestro no tiene la columna 'producto_id'.")

    for raw in it:
        d = {campo: raw[i] for i, campo in mapa.items() if i < len(raw)}
        sku = normalizar_sku(d.get("sku"))
        if not sku or sku == "0":
            continue
        ean = limpiar(d.get("ean"))
        yield {
            "sku": sku,
            "cliente": limpiar(d.get("cliente")) or "VENTURA",
            "descripcion": limpiar(d.get("descripcion"))[:300] or None,
            "ean": ean if ean and ean.upper() not in ("DEFAULT", "0") else None,
            "unidad": limpiar(d.get("unidad"))[:20] or None,
            "familia": limpiar(d.get("familia"))[:60] or None,
            "codigo_alternativo": normalizar_sku(d.get("codigo_alternativo")) or None,
            "inactivo": bool(a_numero(d.get("inactivo"), 0)),
        }
    wb.close()


def filas_a_json(filas: list[FilaPacking]) -> str:
    return json.dumps([f.__dict__ for f in filas], ensure_ascii=False)
