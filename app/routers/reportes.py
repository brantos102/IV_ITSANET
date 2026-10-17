"""Reportes de lectura y Excel con marca, sin alterar datos operativos."""
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.config import BASE_DIR
from app.db import get_db
from app.models import AppRecepcion, Operador
from app.security import requiere_supervisor
from app.services.app_modelo import doc_recepcion

router = APIRouter(prefix='/api/app/reportes', tags=['reportes'])


@router.get('/recepciones')
def recepciones(desde: date | None = None, hasta: date | None = None,
                offset: int = Query(0, ge=0), limite: int = Query(200, ge=1, le=500),
                db: Session = Depends(get_db), _: Operador = Depends(requiere_supervisor)):
    if desde and hasta and desde > hasta:
        raise HTTPException(422, 'La fecha inicial debe ser anterior a la final')
    q = select(AppRecepcion)
    zona = ZoneInfo('America/Guayaquil')
    def utc(d):
        return datetime.combine(d, time.min, zona).astimezone(timezone.utc).replace(tzinfo=None)
    if desde:
        q = q.where(AppRecepcion.creado_en >= utc(desde))
    if hasta:
        q = q.where(AppRecepcion.creado_en < utc(hasta + timedelta(days=1)))
    total = db.scalar(select(func.count()).select_from(q.subquery()))
    filas = db.scalars(q.order_by(AppRecepcion.creado_en.desc(), AppRecepcion.codigo).offset(offset).limit(limite))
    return {'recepciones': [doc_recepcion(r) for r in filas], 'total': total}


class ExcelIn(BaseModel):
    titulo: str = Field(max_length=200)
    filtros: str = Field(default='', max_length=2000)
    resumen: str = Field(default='', max_length=2000)
    columnas: list[str] = Field(min_length=1, max_length=100)
    filas: list[list[str | float | bool | None]] = Field(max_length=20000)


@router.post('/excel')
def excel(datos: ExcelIn, _: Operador = Depends(requiere_supervisor)):
    libro = Workbook(); hoja = libro.active; hoja.title = 'Indicadores'
    logo = Image(BASE_DIR / 'firebase/public/icono.png'); logo.width = 40; logo.height = 40
    hoja.add_image(logo, 'A1')
    hoja['B1'] = 'ITSANET'; hoja['B1'].font = Font(name='Arial', size=20, bold=True, color='1E55D6')
    hoja['A4'] = datos.titulo; hoja['A4'].font = Font(size=14, bold=True)
    hoja['A5'] = datos.filtros; hoja['A6'] = datos.resumen
    def escribir(fila):
        hoja.append(fila)
        for celda in hoja[hoja.max_row]:
            if isinstance(celda.value, str):
                # Los codigos y textos exportados nunca se convierten en formulas.
                celda.data_type = 's'
    escribir(datos.columnas)
    for fila in datos.filas:
        if len(fila) != len(datos.columnas) or any(isinstance(v, str) and len(v) > 32000 for v in fila):
            raise HTTPException(422, 'Fila fuera de los limites del reporte')
        escribir(fila)
    for celda in hoja[7]:
        celda.font = Font(bold=True, color='FFFFFF'); celda.fill = PatternFill('solid', fgColor='1E55D6')
        hoja.column_dimensions[celda.column_letter].width = 22
    hoja.freeze_panes = 'A8'; hoja.auto_filter.ref = f'A7:{hoja.cell(hoja.max_row,len(datos.columnas)).coordinate}'
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.page_setup.orientation = 'landscape'; hoja.page_setup.fitToWidth = 1; hoja.page_setup.fitToHeight = 0
    hoja.print_title_rows = '1:7'
    salida = BytesIO(); libro.save(salida)
    return Response(salida.getvalue(), media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': 'attachment; filename="indicadores-itsanet.xlsx"', 'Cache-Control': 'private, no-store'})
